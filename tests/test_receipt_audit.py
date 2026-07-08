import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from receipt_audit import audit_workbook


class ReceiptAuditTests(unittest.TestCase):
    def test_marks_acquired_and_missing_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            workbook_path = folder / "charges.xlsx"
            output_path = folder / "audited.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.append(["Date", "Description", "Amount"])
            ws.append(["06/01/2026", "Hotel Example", "125.40"])
            ws.append(["06/02/2026", "No Receipt Vendor", "9.99"])
            wb.save(workbook_path)

            (folder / "20260601 Hotel Example 125.40.jpg").write_text("receipt", encoding="utf-8")

            acquired, missing, renamed, moved = audit_workbook(
                excel_path=workbook_path,
                receipts_dir=folder,
                output_path=output_path,
                recursive=True,
                threshold=0.55,
                date_window_days=3,
                rename_receipts=False,
            )

            self.assertEqual(acquired, 1)
            self.assertEqual(missing, 1)
            self.assertEqual(renamed, 0)
            self.assertEqual(moved, 0)

            audited = load_workbook(output_path)
            ws = audited.active
            self.assertEqual(ws["D2"].value, "Acquired")
            self.assertEqual(ws["D3"].value, "Missing")
            self.assertIn("Hotel Example", ws["E2"].value)

    def test_finds_transaction_sheet_when_first_sheet_is_lookup(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            workbook_path = folder / "charges.xlsx"
            output_path = folder / "audited.xlsx"

            wb = Workbook()
            lookup = wb.active
            lookup.title = "GL Accounts"
            lookup.append(["*account_code", "description"])
            lookup.append(["50.21.05", "Direct Travel Cost"])

            ws = wb.create_sheet("Transaction Details")
            ws.append(["Prepared for", "Example"])
            ws.append([])
            ws.append(["Date", "MONTH", "Receipt", "TYPE OF CHARGE", "FREQUENCY", "VENDOR", "Description", "Amount"])
            ws.append(["06/03/2026", "JUNE 2026", None, None, None, "GTLAW", "GTLAW SAMPLE", -5000])
            wb.save(workbook_path)

            (folder / "2026_0603_GTLaw_#8106_$5,000.00.pdf").write_text("receipt", encoding="utf-8")

            acquired, missing, renamed, moved = audit_workbook(
                excel_path=workbook_path,
                receipts_dir=folder,
                output_path=output_path,
                recursive=True,
                threshold=0.55,
                date_window_days=3,
                rename_receipts=False,
            )

            self.assertEqual(acquired, 1)
            self.assertEqual(missing, 0)
            self.assertEqual(renamed, 0)
            self.assertEqual(moved, 0)

            audited = load_workbook(output_path)
            ws = audited["Transaction Details"]
            self.assertEqual(ws["I4"].value, "Acquired")

    def test_prefers_year_sheet_from_workbook_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "03 March 2026"
            folder.mkdir()
            workbook_path = folder / "03 March 2026 AMEX Transaction Details.xlsx"
            output_path = folder / "audited.xlsx"

            wb = Workbook()
            ws_2025 = wb.active
            ws_2025.title = "2025"
            ws_2025.append(["Date", "Vendor", "Amount"])
            ws_2025.append(["03/03/2025", "OLD VENDOR", "1.00"])

            ws_2026 = wb.create_sheet("2026")
            ws_2026.append(["Date", "Vendor", "Amount"])
            ws_2026.append(["03/03/2026", "NEW VENDOR", "2.00"])
            wb.save(workbook_path)

            (folder / "2026_0303_NEW VENDOR_#1234_$2.00.pdf").write_text("receipt", encoding="utf-8")

            acquired, missing, renamed, moved = audit_workbook(
                excel_path=workbook_path,
                receipts_dir=folder,
                output_path=output_path,
                recursive=True,
                threshold=0.55,
                date_window_days=3,
                rename_receipts=False,
            )

            self.assertEqual(acquired, 1)
            self.assertEqual(missing, 0)
            self.assertEqual(renamed, 0)
            self.assertEqual(moved, 0)

            audited = load_workbook(output_path)
            self.assertEqual(audited["2026"]["D2"].value, "Acquired")
            self.assertIsNone(audited["2025"]["D2"].value)

    def test_renames_receipt_to_vendor_invoice_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            workbook_path = folder / "charges.xlsx"
            output_path = folder / "audited.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.append(["Date", "Vendor", "Amount"])
            ws.append(["06/03/2026", "GTLAW", "5000.00"])
            wb.save(workbook_path)

            receipt_path = folder / "2026_0603_GTLaw_#8106_$5,000.00.pdf"
            receipt_path.write_text("receipt", encoding="utf-8")

            acquired, missing, renamed, moved = audit_workbook(
                excel_path=workbook_path,
                receipts_dir=folder,
                output_path=output_path,
                recursive=True,
                threshold=0.55,
                date_window_days=3,
                rename_receipts=True,
            )

            self.assertEqual(acquired, 1)
            self.assertEqual(missing, 0)
            self.assertEqual(renamed, 1)
            self.assertEqual(moved, 0)
            renamed_path = folder / "GTLAW - 8106 - 2026-06-03.pdf"
            self.assertTrue(renamed_path.exists())

            audited = load_workbook(output_path)
            ws = audited.active
            self.assertEqual(ws["E2"].value, "GTLAW - 8106 - 2026-06-03.pdf")

    def test_marks_receipt_column_and_moves_to_done(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            workbook_path = folder / "charges.xlsx"
            output_path = folder / "audited.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.append(["Date", "Receipt", "Vendor", "Amount"])
            ws.append(["06/03/2026", "", "GTLAW", "5000.00"])
            ws.append(["06/04/2026", "yes", "Missing Vendor", "10.00"])
            wb.save(workbook_path)

            receipt_path = folder / "2026_0603_GTLaw_#8106_$5,000.00.pdf"
            receipt_path.write_text("receipt", encoding="utf-8")

            acquired, missing, renamed, moved = audit_workbook(
                excel_path=workbook_path,
                receipts_dir=folder,
                output_path=output_path,
                recursive=True,
                threshold=0.55,
                date_window_days=3,
                rename_receipts=True,
                move_to_done=True,
            )

            self.assertEqual(acquired, 1)
            self.assertEqual(missing, 1)
            self.assertEqual(renamed, 1)
            self.assertEqual(moved, 1)
            self.assertFalse(receipt_path.exists())

            done_receipt = folder / "DONE" / "GTLAW - 8106 - 2026-06-03.pdf"
            self.assertTrue(done_receipt.exists())

            audited = load_workbook(output_path)
            ws = audited.active
            self.assertEqual(ws["B2"].value, "yes")
            self.assertIsNone(ws["B3"].value)
            self.assertEqual(ws["F2"].value, "DONE\\GTLAW - 8106 - 2026-06-03.pdf")


if __name__ == "__main__":
    unittest.main()
