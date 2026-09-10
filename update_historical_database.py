#!/usr/bin/env python3
"""Promote reviewed best-guess vendors into the historical database.

This closes the loop Thomas asked about: once a month's audited workbook has
been reviewed and is considered "final," run this script against it. Any row
that got a category-based best guess (shown with a yellow "Coding Note" flag)
gets folded into historical_vendor_database.xlsx -- so next month, that same
vendor is a known match instead of another guess.

Only best-guess rows are promoted. Confident matches (official reference or
already-in-history) don't need to be added again, and fuzzy matches are
skipped on purpose -- they already matched something close, so re-adding them
as a brand new row would recreate the near-duplicate problem this is meant to
avoid.

Usage:
    python update_historical_database.py "path\\to\\July_receipt_audit.xlsx"
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

CODING_NOTE_REVIEW_TEXT = "Best guess (new vendor) - please review"
HISTORICAL_WORKBOOK = Path(__file__).resolve().parent / "historical_vendor_database.xlsx"
HISTORICAL_SHEET = "Historical Vendor Database"


def find_header_row_and_cols(ws, max_scan=10):
    """Locate the header row and the Vendor/Account Code/Description/Coding
    Note columns by name, wherever they happen to sit in the workbook."""
    for r in range(1, max_scan + 1):
        headers = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v:
                headers[str(v).strip().upper()] = c
        if "VENDOR" in headers and "ACCOUNT CODE" in headers:
            return (
                r,
                headers.get("VENDOR"),
                headers.get("ACCOUNT CODE"),
                headers.get("ACCOUNT DESCRIPTION"),
                headers.get("CODING NOTE"),
            )
    return None, None, None, None, None


def collect_best_guesses(audited_path: Path):
    wb = load_workbook(audited_path, data_only=True)
    # Use whichever sheet actually has the right columns.
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row, vendor_col, code_col, desc_col, note_col = find_header_row_and_cols(ws)
        if not header_row or not note_col:
            continue
        found = []
        for row in range(header_row + 1, ws.max_row + 1):
            note = ws.cell(row=row, column=note_col).value
            vendor = ws.cell(row=row, column=vendor_col).value
            code = ws.cell(row=row, column=code_col).value
            desc = ws.cell(row=row, column=desc_col).value if desc_col else ""
            if note and str(note).strip() == CODING_NOTE_REVIEW_TEXT and vendor and code:
                found.append((str(vendor).strip(), str(code).strip(), str(desc or "").strip()))
        if found or note_col:
            return found
    return []


def load_existing_historical(path: Path):
    entries = []
    existing_vendors = set()
    if path.exists():
        wb = load_workbook(path, data_only=True)
        if HISTORICAL_SHEET in wb.sheetnames:
            ws = wb[HISTORICAL_SHEET]
            for row in range(2, ws.max_row + 1):
                vendor = ws.cell(row=row, column=1).value
                code = ws.cell(row=row, column=2).value
                desc = ws.cell(row=row, column=3).value
                seen = ws.cell(row=row, column=4).value
                consistent = ws.cell(row=row, column=5).value
                if not vendor:
                    continue
                entries.append([vendor, code, desc, seen, consistent])
                existing_vendors.add(str(vendor).strip().upper())
    return entries, existing_vendors


def save_historical(path: Path, entries):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = HISTORICAL_SHEET
    ws.append(["Vendor", "Account Code", "Account Description", "Times Seen", "Consistent"])
    for e in entries:
        ws.append(e)
    wb.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "audited_workbook", type=Path, nargs="?",
        help="Path to a reviewed *_receipt_audit.xlsx file.",
    )
    parser.add_argument(
        "--historical-db",
        type=Path,
        default=HISTORICAL_WORKBOOK,
        help="Path to historical_vendor_database.xlsx (defaults to the copy next to this script).",
    )
    parser.add_argument(
        "--gui",
        action="store_true",
        help="Open a file-picker window instead of typing a path.",
    )
    args = parser.parse_args(argv)

    if args.gui:
        try:
            import tkinter as tk
            from tkinter import filedialog, messagebox
        except Exception as exc:
            print(f"GUI is unavailable: {exc}", file=sys.stderr)
            return 2
        root = tk.Tk()
        root.withdraw()
        picked = filedialog.askopenfilename(
            title="Choose the audited workbook (e.g. *_receipt_audit.xlsx)",
            filetypes=[("Excel workbooks", "*.xlsx *.xlsm *.xls"), ("All files", "*.*")],
        )
        if not picked:
            return 1
        args.audited_workbook = Path(picked)

    if not args.audited_workbook:
        print("An audited workbook path is required unless --gui is used.", file=sys.stderr)
        return 2

    if not args.audited_workbook.exists():
        print(f"File not found: {args.audited_workbook}", file=sys.stderr)
        return 2

    best_guesses = collect_best_guesses(args.audited_workbook)
    if not best_guesses:
        print("No flagged best-guess rows found in this file -- nothing to promote.")
        return 0

    entries, existing_vendors = load_existing_historical(args.historical_db)

    added = []
    for vendor, code, desc in best_guesses:
        key = vendor.upper()
        if key in existing_vendors:
            continue  # already known, skip (avoids creating a duplicate row)
        entries.append([vendor, code, desc, 1, "Promoted from review"])
        existing_vendors.add(key)
        added.append((vendor, code, desc))

    if not added:
        print("All best-guess vendors in this file are already in the historical database. Nothing new to add.")
        return 0

    save_historical(args.historical_db, entries)

    print(f"Promoted {len(added)} vendor(s) into {args.historical_db.name}:")
    for vendor, code, desc in added:
        print(f"  {vendor} -> {code} | {desc}")
    print()
    print("These vendors will now be recognized automatically next time, no more guessing.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
