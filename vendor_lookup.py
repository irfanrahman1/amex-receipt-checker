"""Vendor name extraction and GL account code lookup.

Implements the steps Thomas asked for:
  1. extract_vendor(description) -> pulls a clean vendor name out of the raw
     AmEx "Description" text (e.g. "MSBILL.INFO" -> "MICROSOFT").
  2. VendorLookup.match(vendor) -> looks up that vendor name against two
     sources, in order:
       - vendor_lookup_reference.xlsx: Thomas's official, curated list.
       - historical_vendor_database.xlsx: built from ~17 months of past
         AmEx files. If the company has bought from this vendor before,
         that counts as a known match (no review flag needed), per
         Thomas's guidance.
  3. VendorLookup.match_category(amex_category) -> for vendors found in
     neither source, fall back to a best guess based on AmEx's own
     transaction Category field -- but only for categories that were highly
     consistent (>=85%) across the historical data. This is the ~15% of
     cases Thomas expects will still need a flagged, human-reviewed guess.

All three reference files are meant to keep growing over time -- this
module just reads whatever rows currently exist in them.
"""

import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

# Location of the reference workbooks: same folder as this file.
REFERENCE_WORKBOOK = Path(__file__).resolve().parent / "vendor_lookup_reference.xlsx"
REFERENCE_SHEET = "Lookup Tab Database"
REFERENCE_VENDOR_COL = 2  # column B: Vendor
REFERENCE_CODE_COL = 3  # column C: Account Code
REFERENCE_DESC_COL = 4  # column D: Account Description
REFERENCE_START_ROW = 3  # first data row (row 2 is the header)

# Historical database: built from ~17 months of past AmEx files. Any vendor
# found here has bought-before precedent, so per Thomas it's treated as a
# known match -- no review flag needed, same as the official reference file.
HISTORICAL_WORKBOOK = Path(__file__).resolve().parent / "historical_vendor_database.xlsx"
HISTORICAL_SHEET = "Historical Vendor Database"
HISTORICAL_VENDOR_COL = 1
HISTORICAL_CODE_COL = 2
HISTORICAL_DESC_COL = 3
HISTORICAL_START_ROW = 2

# Category suggestions: for vendors never seen anywhere (not in the official
# list, not in 17 months of history), fall back to a best guess based on
# AmEx's own transaction Category field -- but ONLY for categories that were
# highly consistent (>=85%) across the historical data. This is the ~15%
# Thomas expects to still need a flagged, reviewed guess.
CATEGORY_WORKBOOK = Path(__file__).resolve().parent / "category_suggestions.xlsx"
CATEGORY_SHEET = "Category Suggestions"
CATEGORY_NAME_COL = 1
CATEGORY_CODE_COL = 2
CATEGORY_DESC_COL = 3
CATEGORY_START_ROW = 2

FUZZY_MATCH_THRESHOLD = 0.72

# Known merchant aliases: regex pattern (matched against the uppercased raw
# AmEx description) -> clean vendor name. These handle cases where the raw
# text doesn't clearly spell out who the vendor is (e.g. Microsoft billing
# shows up as "MSBILL.INFO").
VENDOR_ALIASES: List[Tuple[str, str]] = [
    (r"MSBILL\.INFO|A1010BUSD01|MICROSOFT", "MICROSOFT"),
    (r"NYCT PAYGO|OMNY", "MTA"),
    (r"APPLE\.COM", "APPLE"),
    (r"GOOGLE \*YOUTUBE", "YOUTUBE"),
    (r"AMAZON WEB SERVICES|\bAWS\b", "AWS"),
    (r"PAYFLOW/PAYPAL|\bPAYPAL\b", "PAYPAL"),
    (r"SQSP\*|SQUARESPACE", "SQUARESPACE"),
    (r"SMK\*SURVEYMONKEY", "SURVEYMONKEY"),
    (r"WL \*VUE\*TESTING EXAM", "PEARSON VUE"),
    (r"CBI\*ONEIDENTITY|ONEIDENTITY", "ONEIDENTITY"),
    (r"DROPBOX", "DROPBOX"),
    (r"VZWRLSS|VERIZON", "VERIZON"),
    (r"NYTIMES|NY TIMES", "NYTIMES"),
    (r"SEAMLSS\*LITTLEBEET", "LITTLE BEET (SEAMLESS)"),
    (r"SPO\*LABONNESOUPE", "LA BONNE SOUPE"),
    (r"SCALEDAGILE|SCALED AGILE", "SCALED AGILE"),
    (r"IHG POINTS AND CASH|IHG\b", "IHG"),
    (r"\bUBER\b", "UBER"),
    (r"FINANCIAL TIMES", "FINANCIAL TIMES"),
    (r"GODADDY", "GODADDY"),
    (r"CANVA", "CANVA"),
    (r"IN \*K AND GE LLC", "K AND GE LLC"),
    (r"UPWORK", "UPWORK"),
    (r"ZOOM\.COM|\bZOOM\b", "ZOOM"),
    (r"BLS\*CATS", "CATSONE"),
    (r"LINKEDIN", "LINKEDIN"),
    (r"OPENAI|CHATGPT", "CHATGPT"),
    (r"X CORP", "X CORP (TWITTER)"),
    (r"GONITRO", "GONITRO"),
    (r"GSA IFF REMITTANCE", "GSA REMITTANCE"),
    (r"AUTOPAY PAYMENT", "AUTOPAY PAYMENT (INTERNAL)"),
    (r"BAMBOOHR|\bBAMBOO\b", "BAMBOO"),
    (r"CONSTANT CONTACT", "CONSTANT CONTACT"),
    (r"VIVA\*", "VIVA WALLET"),
    (r"AUTOHELLAS|HERTZ", "HERTZ"),
    (r"GOVTRIBE", "GOVTRIBE"),
    (r"SP PLUS CORP|SP PLUS PARKING", "SP PLUS PARKING"),
    (r"WIDEPOINT", "WIDEPOINT"),
    (r"UNITED AIRLINES", "UNITED AIRLINES - FLIGHT"),
    (r"MACYS", "MACYS"),
    (r"GIFTOGRAM", "GIFTOGRAM"),
    (r"VONAGE", "VONAGE"),
    (r"BETTER BUSINESS BURE", "BETTER BUSINESS BUREAU"),
    (r"CRAINS NEWYORK BUSIN|\bCRAINS\b", "CRAINS"),
]

# Known junk prefixes that precede the real vendor name in AmEx descriptions.
PREFIX_STRIP = [
    r"^AplPay\s+", r"^SPO\*", r"^SMK\*", r"^WL \*", r"^CBI\*",
    r"^IN \*", r"^BLS\*", r"^SQSP\*\s*",
]


def extract_vendor(description: str) -> str:
    """Pull a clean vendor name out of a raw AmEx description string."""
    if not description:
        return ""
    desc = description.strip()
    upper = desc.upper()

    for pattern, vendor in VENDOR_ALIASES:
        if re.search(pattern, upper):
            return vendor

    cleaned = desc
    for p in PREFIX_STRIP:
        cleaned = re.sub(p, "", cleaned, flags=re.IGNORECASE)

    cleaned = re.sub(r"\s{2,}[A-Za-z .]+\s+[A-Z]{2}\s*$", "", cleaned)
    cleaned = re.sub(r"\s{2,}[A-Z]{2}\s*$", "", cleaned)

    cleaned = re.sub(r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b", "", cleaned)
    cleaned = re.sub(r"\(\d{3}\)\d{3}-\d{4}", "", cleaned)

    cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" .*-")

    return cleaned.upper() if cleaned else desc.upper()


def extract_frequency_key(description: str) -> str:
    """Like extract_vendor(), but for a vendor whose AmEx lines mix several
    genuinely different recurring products under one generic name -- so
    Frequency classification needs a finer-grained key than account-code
    lookup does.

    Thomas confirmed at the Sept 28 check-in that Microsoft is this kind of
    case: unlike Apple, Microsoft's billing descriptions usually carry real
    product detail (which plan, which service), so there's likely a clean
    monthly/annual pattern *per product* that extract_vendor()'s blanket
    "MICROSOFT" bucket was hiding -- same root cause as the Apple problem,
    just solvable here because the raw text has more to work with. Apple is
    deliberately NOT split this way: Thomas said its descriptions don't
    carry enough detail to make it worth doing, and it's lowest priority
    regardless ("kind of a black box... we'd do that absolutely last").

    IMPORTANT: this key is used ONLY for frequency classification
    (frequency_lookup.py / build_frequency_database.py). It must never
    replace extract_vendor()'s output for account-code lookup --
    historical_vendor_database.xlsx and vendor_lookup_reference.xlsx are
    keyed on the collapsed "MICROSOFT" name Thomas already coded, and
    splitting that for account coding too would turn previously-confident
    Microsoft matches into new "needs review" rows for no reason.
    """
    if not description:
        return ""
    desc = description.strip()
    upper = desc.upper()

    # Opaque reseller billing codes carry no real product detail to keep,
    # so there's nothing finer to extract -- same collapsed bucket as before.
    if re.search(r"MSBILL\.INFO|A1010BUSD01", upper):
        return "MICROSOFT"

    if "MICROSOFT" in upper:
        cleaned = desc
        for p in PREFIX_STRIP:
            cleaned = re.sub(p, "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\s{2,}[A-Za-z .]+\s+[A-Z]{2}\s*$", "", cleaned)
        cleaned = re.sub(r"\s{2,}[A-Z]{2}\s*$", "", cleaned)
        cleaned = re.sub(r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b", "", cleaned)
        cleaned = re.sub(r"\(\d{3}\)\d{3}-\d{4}", "", cleaned)
        cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" .*-")
        return cleaned.upper() if cleaned else "MICROSOFT"

    # Everything else: no known mixed-billing problem, same key as coding.
    return extract_vendor(description)


def _normalize(name: str) -> str:
    n = name.upper()
    n = re.sub(r"[^A-Z0-9 &]", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    return n


@dataclass
class VendorLookupResult:
    vendor_input: str
    matched_vendor: Optional[str]
    account_code: Optional[str]
    account_description: Optional[str]
    confidence: str  # "exact", "fuzzy", "historical", "category_guess", "not_found"


class VendorLookup:
    """Looks up a vendor name, checking sources in order:
    1. vendor_lookup_reference.xlsx -- Thomas's official, curated list.
    2. historical_vendor_database.xlsx -- built from ~17 months of past
       AmEx files. If the company has bought from this vendor before, that
       counts as a known match (no review flag), per Thomas's guidance.
    Vendors found in neither source can still get a category-based best
    guess via match_category() -- see that method for details.
    """

    def __init__(
        self,
        reference_path: Path = REFERENCE_WORKBOOK,
        historical_path: Path = HISTORICAL_WORKBOOK,
        category_path: Path = CATEGORY_WORKBOOK,
        fuzzy_threshold: float = FUZZY_MATCH_THRESHOLD,
    ):
        self.reference_path = reference_path
        self.historical_path = historical_path
        self.category_path = category_path
        self.fuzzy_threshold = fuzzy_threshold
        self.entries: List[Tuple[str, str, str]] = []
        self._exact_index: Dict[str, Tuple[str, str, str]] = {}
        self.historical_entries: List[Tuple[str, str, str]] = []
        self._historical_index: Dict[str, Tuple[str, str, str]] = {}
        self.category_index: Dict[str, Tuple[str, str]] = {}
        self._load()
        self._load_historical()
        self._load_categories()

    def _load(self):
        if not self.reference_path.exists():
            return
        wb = load_workbook(self.reference_path, data_only=True)
        if REFERENCE_SHEET not in wb.sheetnames:
            return
        ws = wb[REFERENCE_SHEET]
        for row in range(REFERENCE_START_ROW, ws.max_row + 1):
            vendor = ws.cell(row=row, column=REFERENCE_VENDOR_COL).value
            code = ws.cell(row=row, column=REFERENCE_CODE_COL).value
            desc = ws.cell(row=row, column=REFERENCE_DESC_COL).value
            if not vendor:
                continue
            vendor = str(vendor).strip()
            code = str(code).strip() if code else ""
            desc = str(desc).strip() if desc else ""
            self.entries.append((vendor, code, desc))
            self._exact_index[_normalize(vendor)] = (vendor, code, desc)

    def _load_historical(self):
        if not self.historical_path.exists():
            return
        wb = load_workbook(self.historical_path, data_only=True)
        if HISTORICAL_SHEET not in wb.sheetnames:
            return
        ws = wb[HISTORICAL_SHEET]
        for row in range(HISTORICAL_START_ROW, ws.max_row + 1):
            vendor = ws.cell(row=row, column=HISTORICAL_VENDOR_COL).value
            code = ws.cell(row=row, column=HISTORICAL_CODE_COL).value
            desc = ws.cell(row=row, column=HISTORICAL_DESC_COL).value
            if not vendor:
                continue
            vendor = str(vendor).strip()
            code = str(code).strip() if code else ""
            desc = str(desc).strip() if desc else ""
            self.historical_entries.append((vendor, code, desc))
            self._historical_index[_normalize(vendor)] = (vendor, code, desc)

    def _load_categories(self):
        if not self.category_path.exists():
            return
        wb = load_workbook(self.category_path, data_only=True)
        if CATEGORY_SHEET not in wb.sheetnames:
            return
        ws = wb[CATEGORY_SHEET]
        for row in range(CATEGORY_START_ROW, ws.max_row + 1):
            cat = ws.cell(row=row, column=CATEGORY_NAME_COL).value
            code = ws.cell(row=row, column=CATEGORY_CODE_COL).value
            desc = ws.cell(row=row, column=CATEGORY_DESC_COL).value
            if not cat or not code:
                continue
            self.category_index[str(cat).strip().upper()] = (str(code).strip(), str(desc).strip())

    def match(self, vendor: str) -> VendorLookupResult:
        if not vendor:
            return VendorLookupResult(vendor, None, None, None, "not_found")

        norm = _normalize(vendor)

        # Tier 1: official reference (exact, then fuzzy)
        hit = self._exact_index.get(norm)
        if hit:
            return VendorLookupResult(vendor, hit[0], hit[1], hit[2], "exact")

        best_score = 0.0
        best_entry = None
        for entry_vendor, code, desc in self.entries:
            score = SequenceMatcher(None, norm, _normalize(entry_vendor)).ratio()
            if score > best_score:
                best_score = score
                best_entry = (entry_vendor, code, desc)
        if best_entry and best_score >= self.fuzzy_threshold:
            return VendorLookupResult(vendor, best_entry[0], best_entry[1], best_entry[2], "fuzzy")

        # Tier 2: historical database (bought before = known, no review flag)
        hist_hit = self._historical_index.get(norm)
        if hist_hit:
            return VendorLookupResult(vendor, hist_hit[0], hist_hit[1], hist_hit[2], "historical")

        best_score = 0.0
        best_entry = None
        for entry_vendor, code, desc in self.historical_entries:
            score = SequenceMatcher(None, norm, _normalize(entry_vendor)).ratio()
            if score > best_score:
                best_score = score
                best_entry = (entry_vendor, code, desc)
        if best_entry and best_score >= self.fuzzy_threshold:
            return VendorLookupResult(vendor, best_entry[0], best_entry[1], best_entry[2], "historical")

        # Tier 3: genuinely never seen before (caller may try match_category next)
        return VendorLookupResult(vendor, None, None, None, "not_found")

    def match_category(self, amex_category: str) -> Optional[Tuple[str, str]]:
        """Best-guess a code from AmEx's own Category field. Only returns a
        result for categories that were highly consistent (>=85%) across 17
        months of historical data. Caller should flag this as needing review.
        """
        if not amex_category:
            return None
        return self.category_index.get(str(amex_category).strip().upper())