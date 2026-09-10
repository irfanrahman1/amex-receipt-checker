#!/usr/bin/env python3
"""Audit an AmEx charge workbook against receipt files in the same folder."""

import argparse
import csv
import datetime as dt
import logging
import re
import sys
from io import BytesIO
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from vendor_lookup import extract_vendor, VendorLookup

try:
    from pypdf import PdfReader, PdfWriter
except Exception:  # pragma: no cover - optional dependency at runtime
    PdfReader = None
    PdfWriter = None

try:
    from PIL import Image, ImageDraw, ImageFont
except Exception:  # pragma: no cover - optional dependency at runtime
    Image = None
    ImageDraw = None
    ImageFont = None

try:
    import pytesseract
except Exception:  # pragma: no cover - optional dependency at runtime
    pytesseract = None

logging.getLogger("pypdf").setLevel(logging.ERROR)


RECEIPT_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".heic",
    ".webp",
    ".tif",
    ".tiff",
}

HEADER_ALIASES = {
    "date": ("date", "transaction date", "posted date", "charge date"),
    "merchant": (
        "merchant",
        "description",
        "vendor",
        "supplier",
        "payee",
        "name",
        "doing business as",
    ),
    "amount": ("amount", "charge", "charge amount", "debit", "total", "extended amount"),
    "category": ("category",),
}

STATUS_HEADER = "Receipt Status"
MATCH_HEADER = "Matched Receipt"
NOTES_HEADER = "Receipt Match Notes"
RECEIPT_FLAG_HEADER = "Receipt"
ACQUIRED_STATUS = "Acquired"
NEEDS_REVIEW_STATUS = "Needs Review"
MISSING_STATUS = "Missing"
STATUS_COLORS = {
    ACQUIRED_STATUS: "C6EFCE",
    NEEDS_REVIEW_STATUS: "FFEB9C",
    MISSING_STATUS: "FFC7CE",
}
VENDOR_HEADER = "Vendor"
ACCOUNT_CODE_HEADER = "Account Code"
ACCOUNT_DESC_HEADER = "Account Description"
CODING_NOTE_HEADER = "Coding Note"
CODING_NOTE_REVIEW_TEXT = "Best guess (new vendor) - please review"
CODING_NOTE_COLOR = "FFEB9C"

@dataclass
class Charge:
    row: int
    date: Optional[dt.date]
    merchant: str
    amount: Optional[Decimal]
    category: str = ""


@dataclass
class Receipt:
    path: Path
    searchable_text: str
    dates: List[dt.date]
    amounts: Set[Decimal]
    filename_amounts: Set[Decimal]


@dataclass
class Match:
    receipt: Optional[Receipt]
    score: float
    notes: str
    status: str = MISSING_STATUS


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).lower().replace(",", "")
    text = re.sub(r"[^a-z0-9.]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_amount(value: object) -> Optional[Decimal]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).copy_abs().quantize(Decimal("0.01"))

    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    negative = "(" in text and ")" in text
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text or text in {"-", "."}:
        return None
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    if negative:
        amount *= Decimal("-1")
    return amount.copy_abs().quantize(Decimal("0.01"))


def parse_date(value: object) -> Optional[dt.date]:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value

    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y", "%m.%d.%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def date_tokens(value: Optional[dt.date]) -> List[str]:
    if value is None:
        return []
    return [
        value.strftime("%m%d%Y"),
        value.strftime("%m%d%y"),
        value.strftime("%Y%m%d"),
        value.strftime("%Y %m%d"),
        value.strftime("%Y %m %d"),
        value.strftime("%m %d %Y").lower(),
        value.strftime("%b %d %Y").lower(),
        value.strftime("%B %d %Y").lower(),
    ]


def extract_dates(text: str) -> List[dt.date]:
    dates: Set[dt.date] = set()
    for year, month, day in re.findall(r"\b(20\d{2})\s*(\d{2})(\d{2})\b", text):
        try:
            dates.add(dt.date(int(year), int(month), int(day)))
        except ValueError:
            pass
    for month, day, year in re.findall(r"\b(\d{1,2})\s+(\d{1,2})\s+(20\d{2})\b", text):
        try:
            dates.add(dt.date(int(year), int(month), int(day)))
        except ValueError:
            pass
    return sorted(dates)


def extract_amounts(text: str) -> Set[Decimal]:
    amounts: Set[Decimal] = set()
    normalized = text.replace(",", "")
    patterns = (
        r"(?<!\d)\$?\s*(-?\d+\.\d{1,2})(?!\d)",
        r"\((\$?\s*\d+\.\d{1,2})\)",
    )
    for pattern in patterns:
        for match in re.findall(pattern, normalized):
            amount = parse_amount(match)
            if amount is not None:
                amounts.add(amount)
    return amounts


def amount_tokens(value: Optional[Decimal]) -> List[str]:
    if value is None:
        return []
    plain = f"{value:.2f}"
    dollars = plain.split(".")[0]
    trimmed = plain.rstrip("0").rstrip(".")
    tokens = [plain, plain.replace(".", ""), f"${plain}", dollars]
    if trimmed != plain:
        tokens.extend([trimmed, f"${trimmed}"])
    return tokens


def merchant_filename_text(merchant: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", " ", str(merchant or ""))
    words = [word for word in text.split() if len(word) > 1]
    return " ".join(words[:5]) or "Receipt"


def safe_filename_part(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*]+', "", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value[:80] or "Receipt"


def invoice_text(receipt: Receipt) -> str:
    generic_values = {"INVOICE", "NUMBER", "NO", "DATE", "AMOUNT", "TOTAL", "RECEIPT"}
    stem_match = re.search(r"#([a-zA-Z0-9-]{3,})", receipt.path.stem)
    if stem_match:
        value = stem_match.group(1).upper()
        if value not in generic_values:
            return value

    text = receipt.searchable_text
    patterns = (
        r"\binvoice\s*(?:number|no|#)?\s*([a-z0-9-]{3,})\b",
        r"\binv(?:\.|\s)+(?:number|no|#)?\s*([a-z0-9-]{3,})\b",
        r"#\s*([a-z0-9-]{3,})\b",
    )
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            value = match.group(1).upper()
            if value not in generic_values and any(char.isdigit() for char in value):
                return value
    return "NoInvoice"


def receipt_filename(charge: Charge, receipt: Receipt) -> str:
    date_part = charge.date.strftime("%Y-%m-%d") if charge.date else "NoDate"
    merchant_part = safe_filename_part(merchant_filename_text(charge.merchant))
    invoice_part = safe_filename_part(invoice_text(receipt))
    return f"{merchant_part} - {invoice_part} - {date_part}{receipt.path.suffix.lower()}"


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    for index in range(2, 1000):
        candidate = path.with_name(f"{stem}_{index}{suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Could not find an available filename for {path}")


def rename_receipt(charge: Charge, receipt: Receipt, target_dir: Optional[Path] = None, rename: bool = True) -> Path:
    target_dir = target_dir or receipt.path.parent
    target_dir.mkdir(parents=True, exist_ok=True)
    target_name = receipt_filename(charge, receipt) if rename else receipt.path.name
    target_path = target_dir / target_name
    if target_path.resolve() == receipt.path.resolve():
        return receipt.path
    new_path = unique_path(target_path)
    if new_path == receipt.path:
        return receipt.path
    receipt.path.rename(new_path)
    receipt.path = new_path
    return new_path


def merge_receipt_pdfs(
    receipt_paths: List[Path],
    output_path: Path,
    receipt_labels: Optional[List[str]] = None,
) -> Tuple[int, List[Path]]:
    if PdfReader is None or PdfWriter is None:
        raise RuntimeError("pypdf is required to create the merged receipt PDF.")

    writer = PdfWriter()
    merged = 0
    skipped = []
    if receipt_paths:
        labels = receipt_labels or [path.name for path in receipt_paths]
        cover = PdfReader(pdf_text_page("AmEx Receipt Packet", [f"Total matched receipts: {len(receipt_paths)}"]))
        writer.add_page(cover.pages[0])
        toc_lines = [f"{index + 1}. {label}" for index, label in enumerate(labels)]
        for start in range(0, len(toc_lines), 34):
            toc = PdfReader(pdf_text_page("Receipt Table of Contents", toc_lines[start : start + 34]))
            writer.add_page(toc.pages[0])

    for path in receipt_paths:
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                reader = PdfReader(str(path))
                for page in reader.pages:
                    writer.add_page(page)
                merged += 1
            elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"} and Image is not None:
                with Image.open(path) as image:
                    if getattr(image, "is_animated", False):
                        image.seek(0)
                    converted = image.convert("RGB")
                    buffer = BytesIO()
                    converted.save(buffer, format="PDF")
                    buffer.seek(0)
                    reader = PdfReader(buffer)
                    for page in reader.pages:
                        writer.add_page(page)
                    merged += 1
            else:
                skipped.append(path)
        except Exception:
            skipped.append(path)

    if merged:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("wb") as output_file:
            writer.write(output_file)
    elif output_path.exists():
        output_path.unlink()

    return merged, skipped


def merchant_score(merchant: str, text: str) -> float:
    merchant_clean = clean_text(merchant)
    if not merchant_clean:
        return 0.0
    if merchant_clean in text:
        return 1.0

    merchant_words = [word for word in merchant_clean.split() if len(word) > 2]
    if not merchant_words:
        return 0.0

    text_words = set(text.split())
    word_hits = sum(1 for word in merchant_words if word in text_words)
    word_score = word_hits / len(merchant_words)
    ratio = SequenceMatcher(None, merchant_clean, text[: max(80, len(merchant_clean) * 4)]).ratio()
    return max(word_score, ratio)


def ocr_image_text(path: Path, enabled: bool) -> str:
    if not enabled or Image is None or pytesseract is None:
        return ""
    try:
        with Image.open(path) as image:
            if getattr(image, "is_animated", False):
                image.seek(0)
            return pytesseract.image_to_string(image)
    except Exception:
        return ""


def extract_pdf_text(path: Path, ocr_enabled: bool = False) -> str:
    if PdfReader is None:
        return ""
    text = ""
    try:
        reader = PdfReader(str(path))
        pages = reader.pages[:3]
        text = " ".join(page.extract_text() or "" for page in pages)
    except Exception:
        text = ""
    # Optional OCR for PDFs is intentionally conservative. We avoid a hard dependency
    # on system PDF rendering tools and only OCR image files directly.
    return text


def load_receipts(folder: Path, excel_path: Path, recursive: bool, ocr_enabled: bool = False) -> List[Receipt]:
    pattern = "**/*" if recursive else "*"
    receipts = []
    for path in folder.glob(pattern):
        if not path.is_file():
            continue
        if path.resolve() == excel_path.resolve():
            continue
        if path.suffix.lower() not in RECEIPT_EXTENSIONS:
            continue
        if clean_text(path.stem).endswith("merged receipts"):
            continue
        text = path.stem
        if path.suffix.lower() == ".pdf":
            text = f"{text} {extract_pdf_text(path, ocr_enabled)}"
        elif path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
            text = f"{text} {ocr_image_text(path, ocr_enabled)}"
        searchable_text = clean_text(text)
        filename_amounts = extract_amounts(path.stem)
        receipts.append(
            Receipt(
                path=path,
                searchable_text=searchable_text,
                dates=extract_dates(searchable_text),
                amounts=extract_amounts(text),
                filename_amounts=filename_amounts,
            )
        )
    return receipts


def find_header_row_and_columns(ws) -> Tuple[int, Dict[str, int]]:
    for row in range(1, min(ws.max_row, 15) + 1):
        seen = {}
        for col in range(1, ws.max_column + 1):
            header = clean_text(ws.cell(row=row, column=col).value)
            if not header:
                continue
            for key, aliases in HEADER_ALIASES.items():
                if header in aliases and key not in seen:
                    seen[key] = col
        if "amount" in seen and ("merchant" in seen or "date" in seen):
            return row, seen
    raise ValueError(
        "Could not identify the workbook headers. Expected columns like Date, "
        "Description/Merchant, and Amount/Charge."
    )


def infer_target_year(path: Path) -> Optional[str]:
    match = re.search(r"\b(20\d{2})\b", str(path))
    return match.group(1) if match else None


def find_charge_sheet(wb, target_year: Optional[str] = None):
    errors = []
    preferred_names = ("transaction details", "master_use me", "master")
    target_year_clean = clean_text(target_year) if target_year else ""

    def sheet_priority(sheet):
        title = clean_text(sheet.title)
        if title in preferred_names:
            return 0
        if target_year_clean and title == target_year_clean:
            return 1
        if target_year_clean and target_year_clean in title:
            return 2
        return 3

    worksheets = sorted(
        wb.worksheets,
        key=lambda sheet: (sheet_priority(sheet), wb.worksheets.index(sheet)),
    )
    for ws in worksheets:
        try:
            header_row, columns = find_header_row_and_columns(ws)
            return ws, header_row, columns
        except ValueError as exc:
            errors.append(f"{ws.title}: {exc}")
    raise ValueError("Could not find a transaction worksheet. " + " | ".join(errors))


def ensure_output_columns(ws, header_row: int) -> Dict[str, int]:
    headers = {
        clean_text(ws.cell(row=header_row, column=col).value): col
        for col in range(1, ws.max_column + 1)
    }
    output = {}
    for label in (
        STATUS_HEADER, 
        MATCH_HEADER, 
        NOTES_HEADER,
        VENDOR_HEADER,
        ACCOUNT_CODE_HEADER,
        ACCOUNT_DESC_HEADER,
        CODING_NOTE_HEADER,
    ):
        key = clean_text(label)
        if key in headers:
            output[label] = headers[key]
        else:
            col = ws.max_column + 1
            ws.cell(row=header_row, column=col).value = label
            output[label] = col
            headers[key] = col
    return output


def receipt_flag_column(ws, header_row: int) -> Optional[int]:
    for col in range(1, ws.max_column + 1):
        if clean_text(ws.cell(row=header_row, column=col).value) == clean_text(RECEIPT_FLAG_HEADER):
            return col
    return None


def iter_charges(ws, header_row: int, columns: Dict[str, int]) -> Iterable[Charge]:
    for row in range(header_row + 1, ws.max_row + 1):
        amount = parse_amount(ws.cell(row=row, column=columns["amount"]).value)
        merchant = ""
        if "merchant" in columns:
            merchant = str(ws.cell(row=row, column=columns["merchant"]).value or "")
        date = None
        if "date" in columns:
            date = parse_date(ws.cell(row=row, column=columns["date"]).value)
        category = ""
        if "category" in columns:
            category = str(ws.cell(row=row, column=columns["category"]).value or "")
        if amount is None and not merchant and date is None:
            continue
        yield Charge(row=row, date=date, merchant=merchant, amount=amount, category=category)


def score_receipt(charge: Charge, receipt: Receipt, date_window_days: int) -> Tuple[float, List[str]]:
    text = receipt.searchable_text
    score = 0.0
    notes = []

    amount_hit = charge.amount is not None and charge.amount in receipt.amounts
    if amount_hit:
        score += 0.45
        notes.append("amount")
    elif charge.amount is not None and receipt.filename_amounts:
        return 0.0, ["filename amount mismatch"]

    date_hit = any(token and clean_text(token) in text for token in date_tokens(charge.date))
    if date_hit:
        score += 0.20
        notes.append("date")
    elif charge.date and date_window_days > 0:
        receipt_dates = receipt.dates
        if receipt_dates:
            closest_days = min(abs((receipt_date - charge.date).days) for receipt_date in receipt_dates)
            if closest_days <= date_window_days:
                score += 0.14
                notes.append(f"date within {closest_days} day(s)")

    merch = merchant_score(charge.merchant, text)
    if merch >= 0.55:
        score += min(0.35, merch * 0.35)
        notes.append(f"merchant {merch:.0%}")

    return score, notes


def best_match(
    charge: Charge,
    receipts: List[Receipt],
    threshold: float,
    review_threshold: float,
    date_window_days: int,
    used_receipts: Set[Path],
) -> Match:
    best = Match(receipt=None, score=0.0, notes="")
    strong_candidates = 0
    for receipt in receipts:
        if receipt.path.resolve() in used_receipts:
            continue
        score, notes = score_receipt(charge, receipt, date_window_days)
        if score >= threshold:
            strong_candidates += 1
        if score > best.score:
            best = Match(receipt=receipt, score=score, notes=", ".join(notes))

    if strong_candidates > 1:
        best.notes = f"{best.notes}; duplicate candidate receipts: {strong_candidates}"
        if best.score < 0.85:
            best.status = NEEDS_REVIEW_STATUS
            return best

    if best.receipt is not None and best.score >= threshold:
        best.status = ACQUIRED_STATUS
        return best
    if best.receipt is not None and best.score >= review_threshold:
        best.status = NEEDS_REVIEW_STATUS
        return best
    return Match(receipt=None, score=best.score, notes=best.notes or "no strong match", status=MISSING_STATUS)


def apply_status_style(ws, row: int, status_col: int, status: str):
    color = STATUS_COLORS.get(status)
    if not color:
        return
    cell = ws.cell(row=row, column=status_col)
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")


def write_reports(
    report_path: Path,
    unmatched_path: Path,
    rows: List[Dict[str, str]],
    receipts: List[Receipt],
    used_receipts: Set[Path],
):
    report_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "row",
        "status",
        "date",
        "merchant",
        "amount",
        "matched_receipt",
        "score",
        "notes",
        "vendor",
        "account_code",
        "account_description",
        "coding_note",
    ]
    with report_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    with unmatched_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["receipt"])
        writer.writeheader()
        for receipt in receipts:
            if receipt.path.resolve() not in used_receipts:
                writer.writerow({"receipt": str(receipt.path)})


def pdf_text_page(title: str, lines: List[str]) -> BytesIO:
    if Image is None:
        raise RuntimeError("Pillow is required to create merged PDF cover pages.")
    width, height = 1700, 2200
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 54)
        body_font = ImageFont.truetype("arial.ttf", 32)
    except Exception:
        title_font = body_font = ImageFont.load_default()
    draw.text((90, 90), title, fill=(25, 35, 50), font=title_font)
    y = 190
    for line in lines:
        draw.text((90, y), line[:95], fill=(35, 35, 35), font=body_font)
        y += 52
        if y > height - 100:
            break
    buffer = BytesIO()
    image.save(buffer, format="PDF")
    buffer.seek(0)
    return buffer


def audit_workbook(
    excel_path: Path,
    receipts_dir: Path,
    output_path: Path,
    recursive: bool,
    threshold: float,
    review_threshold: float,
    date_window_days: int,
    rename_receipts: bool,
    move_to_done: bool = False,
    merged_pdf_path: Optional[Path] = None,
    report_path: Optional[Path] = None,
    unmatched_report_path: Optional[Path] = None,
    ocr_enabled: bool = False,
) -> Tuple[int, int, int, int, int, int]:
    wb = load_workbook(excel_path)
    ws, header_row, columns = find_charge_sheet(wb, infer_target_year(excel_path))
    output_columns = ensure_output_columns(ws, header_row)
    flag_column = receipt_flag_column(ws, header_row)
    receipts = load_receipts(receipts_dir, excel_path, recursive, ocr_enabled=ocr_enabled)
    done_dir = receipts_dir / "DONE" if move_to_done else None
    vendor_lookup = VendorLookup()

    acquired = 0
    missing = 0
    needs_review = 0
    renamed = 0
    moved = 0
    matched_receipts = []
    matched_labels = []
    report_rows: List[Dict[str, str]] = []
    used_receipts: Set[Path] = set()
    for charge in iter_charges(ws, header_row, columns):
        match = best_match(charge, receipts, threshold, review_threshold, date_window_days, used_receipts)
        if match.status == ACQUIRED_STATUS and match.receipt:
            acquired += 1
            old_path = match.receipt.path
            if rename_receipts or move_to_done:
                rename_receipt(charge, match.receipt, target_dir=done_dir, rename=rename_receipts)
                if match.receipt.path != old_path:
                    if match.receipt.path.name != old_path.name:
                        renamed += 1
                    if match.receipt.path.parent.resolve() != old_path.parent.resolve():
                        moved += 1
            used_receipts.add(match.receipt.path.resolve())
            matched_receipts.append(match.receipt.path)
            matched_labels.append(
                f"Row {charge.row}: {charge.date or ''} | {charge.merchant} | "
                f"{charge.amount or ''} | {match.receipt.path.name}"
            )
            if flag_column:
                ws.cell(row=charge.row, column=flag_column).value = "yes"
            ws.cell(row=charge.row, column=output_columns[STATUS_HEADER]).value = ACQUIRED_STATUS
            ws.cell(row=charge.row, column=output_columns[MATCH_HEADER]).value = str(
                match.receipt.path.relative_to(receipts_dir)
            )
        elif match.status == NEEDS_REVIEW_STATUS and match.receipt:
            needs_review += 1
            if flag_column:
                ws.cell(row=charge.row, column=flag_column).value = ""
            ws.cell(row=charge.row, column=output_columns[STATUS_HEADER]).value = NEEDS_REVIEW_STATUS
            ws.cell(row=charge.row, column=output_columns[MATCH_HEADER]).value = str(
                match.receipt.path.relative_to(receipts_dir)
            )
        else:
            missing += 1
            if flag_column:
                ws.cell(row=charge.row, column=flag_column).value = ""
            ws.cell(row=charge.row, column=output_columns[STATUS_HEADER]).value = MISSING_STATUS
            ws.cell(row=charge.row, column=output_columns[MATCH_HEADER]).value = ""
        ws.cell(row=charge.row, column=output_columns[NOTES_HEADER]).value = (
            f"score {match.score:.2f}; {match.notes}"
        )
        status = ws.cell(row=charge.row, column=output_columns[STATUS_HEADER]).value
        apply_status_style(ws, charge.row, output_columns[STATUS_HEADER], status)
        vendor_name = extract_vendor(charge.merchant)
        vendor_result = vendor_lookup.match(vendor_name)
        account_code = vendor_result.account_code or ""
        account_desc = vendor_result.account_description or ""
        coding_note = ""
        if vendor_result.confidence == "not_found":
            category_guess = vendor_lookup.match_category(charge.category)
            if category_guess:
                account_code, account_desc = category_guess
                coding_note = CODING_NOTE_REVIEW_TEXT
        ws.cell(row=charge.row, column=output_columns[VENDOR_HEADER]).value = vendor_name
        ws.cell(row=charge.row, column=output_columns[ACCOUNT_CODE_HEADER]).value = account_code
        ws.cell(row=charge.row, column=output_columns[ACCOUNT_DESC_HEADER]).value = account_desc
        ws.cell(row=charge.row, column=output_columns[CODING_NOTE_HEADER]).value = coding_note
        if coding_note:
            note_cell = ws.cell(row=charge.row, column=output_columns[CODING_NOTE_HEADER])
            note_cell.fill = PatternFill(start_color=CODING_NOTE_COLOR, end_color=CODING_NOTE_COLOR, fill_type="solid")
        report_rows.append(
            {
                "row": str(charge.row),
                "status": str(status),
                "date": str(charge.date or ""),
                "merchant": charge.merchant,
                "amount": str(charge.amount or ""),
                "matched_receipt": str(ws.cell(row=charge.row, column=output_columns[MATCH_HEADER]).value or ""),
                "score": f"{match.score:.2f}",
                "notes": match.notes,
                "vendor": vendor_name,
                "account_code": vendor_result.account_code or "",
                "account_description": vendor_result.account_description or "",
            }
        )

    wb.save(output_path)
    merged = 0
    if merged_pdf_path:
        merged, skipped = merge_receipt_pdfs(matched_receipts, merged_pdf_path, matched_labels)
        if skipped:
            skipped_text = ", ".join(str(path) for path in skipped[:5])
            if len(skipped) > 5:
                skipped_text += f", and {len(skipped) - 5} more"
            print(f"Skipped receipts that could not be merged: {skipped_text}", file=sys.stderr)
    if report_path:
        write_reports(
            report_path,
            unmatched_report_path or report_path.with_name(f"{report_path.stem}_unmatched_receipts.csv"),
            report_rows,
            receipts,
            used_receipts,
        )
    return acquired, needs_review, missing, renamed, moved, merged


def default_output_path(excel_path: Path) -> Path:
    return excel_path.with_name(f"{excel_path.stem}_receipt_audit{excel_path.suffix}")


def default_merged_pdf_path(excel_path: Path) -> Path:
    return excel_path.with_name(f"{excel_path.stem}_merged_receipts.pdf")


def default_report_path(excel_path: Path) -> Path:
    return excel_path.with_name(f"{excel_path.stem}_audit_report.csv")


def candidate_workbooks(root: Path) -> List[Path]:
    workbooks = []
    for path in root.rglob("*.xlsx"):
        name = path.name.lower()
        if "receipt_audit" in name or name.startswith("~$"):
            continue
        if "transaction details" in name:
            workbooks.append(path)
    return sorted(workbooks)


def run_single(args, excel_path: Path) -> Tuple[int, int, int, int, int, int]:
    receipts_dir = (args.receipts_dir or excel_path.parent).expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output and not args.batch_root
        else default_output_path(excel_path).resolve()
    )
    merged_pdf_path = None
    if not args.no_merge_pdf:
        merged_pdf_path = (
            args.merged_pdf_output.expanduser().resolve()
            if args.merged_pdf_output and not args.batch_root
            else default_merged_pdf_path(excel_path).resolve()
        )
    report_path = None
    unmatched_path = None
    if not args.no_reports:
        report_path = (
            args.report_output.expanduser().resolve()
            if args.report_output and not args.batch_root
            else default_report_path(excel_path).resolve()
        )
        unmatched_path = report_path.with_name(f"{report_path.stem}_unmatched_receipts.csv")

    acquired, needs_review, missing, renamed, moved, merged = audit_workbook(
        excel_path=excel_path,
        receipts_dir=receipts_dir,
        output_path=output_path,
        recursive=not args.no_recursive,
        threshold=args.threshold,
        review_threshold=args.review_threshold,
        date_window_days=args.date_window_days,
        rename_receipts=args.rename_receipts,
        move_to_done=args.move_to_done,
        merged_pdf_path=merged_pdf_path,
        report_path=report_path,
        unmatched_report_path=unmatched_path,
        ocr_enabled=args.ocr,
    )
    print(f"Saved audited workbook: {output_path}")
    if merged_pdf_path:
        print(f"Saved merged receipt PDF: {merged_pdf_path}")
        print(f"Merged receipts: {merged}")
    if report_path:
        print(f"Saved audit report: {report_path}")
        print(f"Saved unmatched receipts report: {unmatched_path}")
    print(f"Acquired: {acquired}")
    print(f"Needs review: {needs_review}")
    print(f"Missing: {missing}")
    if args.rename_receipts:
        print(f"Renamed receipts: {renamed}")
    if args.move_to_done:
        print(f"Moved to DONE: {moved}")
    return acquired, needs_review, missing, renamed, moved, merged


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Mark AmEx workbook charges as Acquired or Missing based on receipt files."
    )
    parser.add_argument("excel", nargs="?", type=Path, help="Path to the AmEx Excel workbook.")
    parser.add_argument(
        "--receipts-dir",
        type=Path,
        help="Folder containing receipts. Defaults to the workbook's folder.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="Where to save the audited workbook. Defaults to *_receipt_audit.xlsx.",
    )
    parser.add_argument(
        "--no-recursive",
        action="store_true",
        help="Only check files directly in the receipt folder, not subfolders.",
    )
    parser.add_argument(
        "--threshold",
        type=float,
        default=0.55,
        help="Match threshold from 0.0 to 1.0. Default: 0.55.",
    )
    parser.add_argument(
        "--review-threshold",
        type=float,
        default=0.40,
        help="Score from this value to --threshold is marked Needs Review. Default: 0.40.",
    )
    parser.add_argument(
        "--date-window-days",
        type=int,
        default=3,
        help="Allow receipt dates this many days before/after the workbook date. Default: 3.",
    )
    parser.add_argument(
        "--rename-receipts",
        action="store_true",
        help="Rename matched receipt files to Vendor - Invoice - Date.ext.",
    )
    parser.add_argument(
        "--move-to-done",
        action="store_true",
        help="Move matched receipts into a DONE subfolder, matching the finished-folder workflow.",
    )
    parser.add_argument(
        "--merged-pdf-output",
        type=Path,
        help="Where to save the combined PDF of acquired receipts. Defaults to *_merged_receipts.pdf.",
    )
    parser.add_argument(
        "--no-merge-pdf",
        action="store_true",
        help="Do not create a combined PDF of acquired receipts.",
    )
    parser.add_argument(
        "--report-output",
        type=Path,
        help="Where to save the CSV audit report. Defaults to *_audit_report.csv.",
    )
    parser.add_argument(
        "--no-reports",
        action="store_true",
        help="Do not create CSV audit and unmatched receipt reports.",
    )
    parser.add_argument(
        "--ocr",
        action="store_true",
        help="Use OCR for image receipts when pytesseract and Tesseract are installed.",
    )
    parser.add_argument(
        "--batch-root",
        type=Path,
        help="Process every *Transaction Details.xlsx workbook under this root.",
    )
    parser.add_argument(
        "--gui",
        action="store_true",
        help="Open a simple file-picker interface.",
    )
    return parser


def run_gui() -> int:
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox
    except Exception as exc:
        print(f"GUI is unavailable: {exc}", file=sys.stderr)
        return 2

    root = tk.Tk()
    root.withdraw()
    excel = filedialog.askopenfilename(
        title="Choose AmEx workbook",
        filetypes=[("Excel workbooks", "*.xlsx *.xlsm *.xls"), ("All files", "*.*")],
    )
    if not excel:
        return 1
    receipts_dir = filedialog.askdirectory(title="Choose receipt folder")
    if not receipts_dir:
        return 1
    excel_path = Path(excel).resolve()
    receipts_path = Path(receipts_dir).resolve()
    acquired, needs_review, missing, renamed, moved, merged = audit_workbook(
        excel_path=excel_path,
        receipts_dir=receipts_path,
        output_path=default_output_path(excel_path),
        recursive=True,
        threshold=0.55,
        review_threshold=0.40,
        date_window_days=3,
        rename_receipts=True,
        move_to_done=True,
        merged_pdf_path=default_merged_pdf_path(excel_path),
        report_path=default_report_path(excel_path),
    )
    messagebox.showinfo(
        "AmEx Receipt Checker",
        (
            f"Audit complete.\n\n"
            f"Acquired: {acquired}\n"
            f"Needs review: {needs_review}\n"
            f"Missing: {missing}\n"
            f"Renamed: {renamed}\n"
            f"Moved to DONE: {moved}\n"
            f"Merged receipts: {merged}"
        ),
    )
    return 0


def main(argv: Optional[List[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    if args.gui:
        return run_gui()

    if args.batch_root:
        root = args.batch_root.expanduser().resolve()
        if not root.exists():
            print(f"Batch root not found: {root}", file=sys.stderr)
            return 2
        workbooks = candidate_workbooks(root)
        if not workbooks:
            print(f"No transaction workbooks found under: {root}", file=sys.stderr)
            return 2
        total_acquired = total_needs_review = total_missing = total_renamed = total_moved = total_merged = 0
        for workbook in workbooks:
            print(f"\nProcessing: {workbook}")
            acquired, needs_review, missing, renamed, moved, merged = run_single(args, workbook)
            total_acquired += acquired
            total_needs_review += needs_review
            total_missing += missing
            total_renamed += renamed
            total_moved += moved
            total_merged += merged
        print("\nBatch complete")
        print(f"Total acquired: {total_acquired}")
        print(f"Total needs review: {total_needs_review}")
        print(f"Total missing: {total_missing}")
        print(f"Total renamed: {total_renamed}")
        print(f"Total moved to DONE: {total_moved}")
        print(f"Total merged receipts: {total_merged}")
        return 0

    if not args.excel:
        print("Excel workbook path is required unless --batch-root or --gui is used.", file=sys.stderr)
        return 2

    excel_path = args.excel.expanduser().resolve()

    if not excel_path.exists():
        print(f"Excel file not found: {excel_path}", file=sys.stderr)
        return 2
    if args.receipts_dir and not args.receipts_dir.expanduser().resolve().exists():
        print(f"Receipt folder not found: {args.receipts_dir}", file=sys.stderr)
        return 2

    run_single(args, excel_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
